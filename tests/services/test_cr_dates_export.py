"""Customer-return Excel export dates: IST values and IST day filters.

The export used to write str(rtv_date) — the stored naive-UTC clock face, so a
return raised at 02:00 IST on 19 Oct read "2026-10-18 20:30:00". RTV Date and
Created At are now real Excel dates on the IST clock (Excel has no time zones),
formatted so they read as dates and sort and filter as dates. The column names
and their order are unchanged. Uses the fake connection and HTTP client of
test_cr_dates_read.

Run:  PYTHONPATH=. python -m pytest tests/services/test_cr_dates_export.py
"""
from __future__ import annotations

import asyncio
from datetime import datetime, timedelta, timezone
from io import BytesIO

from openpyxl import load_workbook

from app.modules.customer_returns.services import export_xlsx
from app.modules.customer_returns.services import query_service as q
from app.modules.customer_returns.services.query_service import EXPORT_COLUMNS
from tests.services import test_cr_dates_read as rd
from tests.services.test_cr_dates_read import client  # noqa: F401  (the HTTP fixture)

RTV_DATE = EXPORT_COLUMNS.index("RTV Date") + 1
CREATED_AT = EXPORT_COLUMNS.index("Created At") + 1


def _export(conn, **filters) -> list:
    return asyncio.run(q.export_cr_records(conn, company="CFPL", **filters))


def _sheet(rows: list):
    return load_workbook(export_xlsx.build_export_workbook(rows, set())).active


def test_export_rows_carry_ist_datetimes():
    conn = rd.FakeConn(rd._hdr(created_ts=datetime(2026, 10, 18, 20, 30, 5)),
                       lines=[rd._line()], boxes=[rd._box()])
    [row] = _export(conn)
    assert list(row) == EXPORT_COLUMNS                     # same 33 columns, same order
    assert row["RTV Date"].isoformat() == "2026-10-19T02:00:00+05:30"
    assert row["Created At"].isoformat() == "2026-10-19T02:00:05+05:30"
    [undated] = _export(rd.FakeConn(rd._hdr(rtv_date=None, created_ts=None)))
    assert undated["RTV Date"] == "" and undated["Created At"] == ""


def test_the_workbook_writes_ist_dates_as_excel_dates():
    conn = rd.FakeConn(rd._hdr(created_ts=datetime(2026, 10, 18, 20, 30, 5)),
                       lines=[rd._line()], boxes=[rd._box()])
    ws = _sheet(_export(conn))
    assert [c.value for c in ws[1]] == EXPORT_COLUMNS
    rtv_date, created_at = ws.cell(row=2, column=RTV_DATE), ws.cell(row=2, column=CREATED_AT)
    assert rtv_date.value == datetime(2026, 10, 19, 2, 0)            # IST wall clock, naive
    assert created_at.value == datetime(2026, 10, 19, 2, 0, 5)
    assert rtv_date.is_date and created_at.is_date
    assert rtv_date.number_format == "dd-mmm-yyyy"
    assert created_at.number_format == "dd-mmm-yyyy hh:mm"
    for col in (RTV_DATE, CREATED_AT):                    # wide enough not to show ####
        assert ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width >= 18
    # Every other column is written as before.
    assert ws.cell(row=2, column=EXPORT_COLUMNS.index("RTV ID") + 1).value == rd.CR_ID
    assert ws.cell(row=2, column=EXPORT_COLUMNS.index("Box Net Weight") + 1).value == 25.0


def test_any_aware_datetime_is_written_on_the_ist_clock():
    base = {c: "" for c in EXPORT_COLUMNS}
    rows = [dict(base, **{"RTV Date": datetime(2026, 10, 18, 20, 30, tzinfo=timezone.utc),
                          "Created At": datetime(2026, 10, 19, 5, 30, tzinfo=timezone(timedelta(hours=9)))}),
            dict(base, **{"RTV ID": "CR-UNDATED"})]
    ws = _sheet(rows)
    assert ws.cell(row=2, column=RTV_DATE).value == datetime(2026, 10, 19, 2, 0)
    assert ws.cell(row=2, column=CREATED_AT).value == datetime(2026, 10, 19, 2, 0)
    assert ws.cell(row=3, column=RTV_DATE).value in (None, "")          # an empty date stays empty


def test_the_export_route_filters_and_writes_on_ist_days(client):
    conn = rd.FakeConn(rd._hdr(), rd._hdr(id=rd.HID + 1, rtv_id="CR-20261018235900",
                                          rtv_date=datetime(2026, 10, 18, 18, 29)),
                       lines=[rd._line()], boxes=[rd._box()])
    res = client(conn).get("/api/v1/customer-returns/export",
                           params={"company": "CFPL", "from_date": "19-10-2026", "to_date": "19-10-2026"})
    assert res.status_code == 200, res.text
    ws = load_workbook(BytesIO(res.content)).active
    assert [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)] == [rd.CR_ID]
    assert ws.cell(row=2, column=RTV_DATE).value == datetime(2026, 10, 19, 2, 0)
    [(_kind, sql, args)] = [c for c in conn.calls if "LEFT JOIN" in c[1]]
    assert args == (datetime(2026, 10, 18, 18, 30), datetime(2026, 10, 19, 18, 30))
