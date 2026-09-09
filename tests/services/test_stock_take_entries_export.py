"""LIVE-SQL test for the floor count Excel export (GET /entries/export).

Runs the real service and the real workbook builder, then reads the .xlsx back
with openpyxl and asserts the numbers IN THE SHEET against the numbers in SQL.
Checking the service's return value alone would miss the half of this feature
that is the spreadsheet.

`new_stock_entries` is owned by the separate Stock Take app and lives in the AWS
RDS `warehouse_db`; it does not exist in the Supabase database `.env` can be
pointed at, so these statements cannot be validated by a FakeConn. Point the run
at RDS to exercise them:

    STOCKTAKE_TEST_DATABASE_URL=postgresql://.../warehouse_db \
    PYTHONPATH=. .venv/Scripts/python tests/services/test_stock_take_entries_export.py

Otherwise:

    PYTHONPATH=. .venv/Scripts/python tests/services/test_stock_take_entries_export.py

Expected values are read from the database at run time, never hardcoded: people
are still counting, and a pinned snapshot would turn "a floor manager submitted
a batch" into a wall of red. Pure reads -- nothing is written.
"""
import asyncio
import io
import os

import asyncpg
from openpyxl import load_workbook

from app.config import Settings
from app.modules.stock_take.services import entries_export, export_xlsx
from app.modules.stock_take.services import business_day

_passed = 0
_failed = 0

# The IST business day of a new_stock_entries row. Written out rather than
# imported, so that changing business_day breaks this file loudly instead of
# silently moving the expected values with it. The assertions below are what keep
# the copy honest.
#
# ONE step, not two: new_stock_entries.created_at is timestamptz. The two-step
# form belongs to the floor app's stocktake_entries, whose column is NAIVE and
# holds UTC; using it here would shift every business day by -5:30 while still
# parsing and still returning dates.
ED = "(created_at AT TIME ZONE 'Asia/Kolkata')::date"
TBL = "new_stock_entries"
assert ED == business_day.ENTRY_DAY,     "business_day.ENTRY_DAY is now %r -- update ED and re-check every expected value"     % (business_day.ENTRY_DAY,)
assert TBL == business_day.ENTRIES_TABLE,     "business_day.ENTRIES_TABLE is now %r -- this test is reading the wrong table"     % (business_day.ENTRIES_TABLE,)

# What the export considers a floor count: submitted, and not a console
# adjustment written back into this table.
COUNTED = ("(status IS NULL OR status != 'draft')"
           " AND (source_kind IS NULL OR source_kind = 'COUNT')")


def check(label, cond, extra=""):
    global _passed, _failed
    if cond:
        _passed += 1
        print("  PASS  %s" % label)
    else:
        _failed += 1
        print("  FAIL  %s %s" % (label, extra))


def near(a, b):
    """Within a hair -- NUMERIC SUM vs Python float."""
    return abs(float(a) - float(b)) < 0.01


async def main():
    url = os.getenv("STOCKTAKE_TEST_DATABASE_URL") or Settings().DATABASE_URL
    conn = await asyncpg.connect(url)
    try:
        if not await conn.fetchval("SELECT to_regclass('new_stock_entries')"):
            print(
                "\n  SKIP  new_stock_entries is not in the configured database.\n"
                "        This app's DATABASE_URL points somewhere without the Stock Take\n"
                "        tables (the Supabase config has none). Set STOCKTAKE_TEST_DATABASE_URL\n"
                "        to the RDS warehouse_db to exercise these statements.\n"
            )
            return

        # ── [1] Unfiltered: the sheet agrees with SQL ────────────────────
        print("\n[1] Unfiltered export")
        rows, applied, drafts = await entries_export.fetch_entries(conn)
        truth = await conn.fetchrow(f"""
            SELECT COUNT(*)::int                            AS n,
                   COALESCE(SUM(total_weight), 0)           AS kg,
                   COALESCE(SUM(total_quantity), 0)         AS qty,
                   COUNT(*) FILTER (WHERE verified IS TRUE)::int AS verified
              FROM new_stock_entries WHERE {COUNTED}""")
        check("row count matches SQL", len(rows) == truth["n"],
              "got %d want %d" % (len(rows), truth["n"]))
        check("verified count matches SQL",
              sum(1 for r in rows if r.get("verified") is True) == truth["verified"])

        draft_truth = await conn.fetchval(
            "SELECT COUNT(*)::int FROM new_stock_entries WHERE status = 'draft'"
            " AND (source_kind IS NULL OR source_kind = 'COUNT')")
        check("drafts are excluded and counted", drafts == draft_truth,
              "got %d want %d" % (drafts, draft_truth))
        check("no draft row reached the sheet",
              all(r.get("status") != "draft" for r in rows))
        check("no console adjustment reached the sheet",
              await conn.fetchval(
                  "SELECT COUNT(*)::int FROM new_stock_entries"
                  " WHERE source_kind = 'ADJUSTMENT' AND id = ANY($1::int[])",
                  [r["id"] for r in rows]) == 0)

        # ── [2] The workbook itself ──────────────────────────────────────
        print("\n[2] Workbook structure")
        wb = load_workbook(io.BytesIO(export_xlsx.build_entries_workbook(
            rows, applied, drafts, "test@candorfoods.in").getvalue()))
        check("Summary sheet is always present", "Summary" in wb.sheetnames,
              "sheets %r" % (wb.sheetnames,))
        first = next((n for n in wb.sheetnames if n != "Summary"), None)
        check("a stock-type sheet was written", first is not None)

        ws = wb[first]
        header = [ws.cell(row=4, column=i).value for i in range(1, len(export_xlsx.ENTRY_COLUMNS) + 1)]
        # Columns 1-24 are the floor app's own columns in its own order, so the
        # two files can be diffed row-for-row; Batch ID is APPENDED, never
        # inserted, so nothing existing shifts.
        check("header starts at the floor app's first column", header[0] == "Entry ID",
              "got %r" % header[0])
        check("column 24 is still the floor app's last column",
              header[23] == "Verified At (Signature)", "got %r" % header[23])
        check("Batch ID is appended as column 25", header[24] == "Batch ID",
              "got %r" % header[24])
        check("no duplicate headers", len(set(header)) == len(header))

        # ── [3] Totals in the sheet, not just in the service ─────────────
        print("\n[3] Summary sheet arithmetic")
        sm = wb["Summary"]
        gt = next(r for r in range(1, sm.max_row + 1)
                  if sm.cell(row=r, column=1).value == "GRAND TOTAL")
        check("GRAND TOTAL entries matches SQL",
              sm.cell(row=gt, column=2).value == truth["n"],
              "got %r want %d" % (sm.cell(row=gt, column=2).value, truth["n"]))
        check("GRAND TOTAL weight matches SQL",
              near(sm.cell(row=gt, column=4).value, truth["kg"]),
              "got %r want %s" % (sm.cell(row=gt, column=4).value, truth["kg"]))
        check("GRAND TOTAL quantity matches SQL",
              near(sm.cell(row=gt, column=3).value, truth["qty"]))
        # The floor app writes these with toFixed(2), which lands them in Excel
        # as TEXT and cannot be summed. These must be real numbers.
        check("weight cell is a number, not text",
              isinstance(sm.cell(row=gt, column=4).value, (int, float)),
              "got %r" % type(sm.cell(row=gt, column=4).value))
        # Express drops any stock_type outside its two hardcoded buckets from
        # both sheets while still counting it in the grand total. Every bucket
        # gets a row here, so the column always adds up.
        buckets = [r for r in range(5, gt) if isinstance(sm.cell(row=r, column=2).value, int)]
        check("per-bucket rows sum to the grand total",
              sum(sm.cell(row=r, column=2).value for r in buckets) == truth["n"],
              "buckets %d vs total %d"
              % (sum(sm.cell(row=r, column=2).value for r in buckets), truth["n"]))

        # ── [4] The date window is cut on the IST day ────────────────────
        print("\n[4] IST business day, not the server's")
        day = await conn.fetchval(
            f"SELECT MAX({ED}) FROM new_stock_entries WHERE {COUNTED}")
        one_day, _, _ = await entries_export.fetch_entries(
            conn, date_from=day.isoformat(), date_to=day.isoformat())
        ist_n = await conn.fetchval(
            f"SELECT COUNT(*)::int FROM new_stock_entries WHERE {COUNTED} AND {ED} = $1", day)
        check("a one-day export uses the IST window", len(one_day) == ist_n,
              "got %d want %d" % (len(one_day), ist_n))
        # And prove the two rules genuinely differ on this dataset, so the check
        # above is not passing because the correction is a no-op.
        skew = await conn.fetchval(
            f"SELECT COUNT(*)::int FROM new_stock_entries"
            f" WHERE {COUNTED} AND created_at::date <> {ED}")
        check("the naive and IST days really do disagree here", skew > 0,
              "0 rows differ -- this test would pass either way")
        print("        %d row(s) fall on a different day under a UTC cut" % skew)

        # ── [5] Filters narrow both the rows and the draft footnote ──────
        print("\n[5] Filters")
        wh = await conn.fetchval(
            "SELECT UPPER(TRIM(warehouse)) FROM new_stock_entries"
            " WHERE warehouse IS NOT NULL GROUP BY 1 ORDER BY COUNT(*) DESC LIMIT 1")
        f_rows, f_applied, f_drafts = await entries_export.fetch_entries(conn, warehouse=[wh])
        f_n = await conn.fetchval(
            f"SELECT COUNT(*)::int FROM new_stock_entries"
            f" WHERE {COUNTED} AND UPPER(TRIM(warehouse)) = $1", wh)
        check("warehouse filter matches SQL", len(f_rows) == f_n,
              "got %d want %d" % (len(f_rows), f_n))
        check("applied echoes the filter for the sheet header",
              f_applied.get("warehouse") == [wh], "got %r" % f_applied.get("warehouse"))
        # The Express version applies only `warehouse` to its draft count, so a
        # one-floor one-day export is footnoted with the whole warehouse's drafts.
        check("the draft footnote respects the filter too", f_drafts <= drafts,
              "filtered %d > unfiltered %d" % (f_drafts, drafts))

        # ── [6] An empty match is still a readable file ──────────────────
        print("\n[6] Empty match")
        e_rows, e_applied, e_drafts = await entries_export.fetch_entries(
            conn, warehouse=["__NO_SUCH_WAREHOUSE__"])
        check("nothing matched", len(e_rows) == 0)
        e_wb = load_workbook(io.BytesIO(export_xlsx.build_entries_workbook(
            e_rows, e_applied, e_drafts, "test").getvalue()))
        # A workbook with zero sheets is not a valid .xlsx, and a 404 on a
        # download is indistinguishable from a broken endpoint at the browser.
        check("the workbook still opens, with only the Summary",
              e_wb.sheetnames == ["Summary"], "got %r" % (e_wb.sheetnames,))

        # ── [7] Unusable dates are rejected, not ignored ─────────────────
        print("\n[7] Date validation")
        for bad, why in (({"date_from": "05-09-2026"}, "DD-MM-YYYY"),
                         ({"date_to": "not-a-date"}, "not a date"),
                         ({"date_from": "2026-09-05", "date_to": "2026-09-01"}, "reversed range")):
            try:
                await entries_export.fetch_entries(conn, **bad)
                # Dropping it would widen the export to every row ever counted,
                # which is indistinguishable from success in a downloaded file.
                check("rejects %s" % why, False, "accepted %r" % bad)
            except ValueError:
                check("rejects %s" % why, True)

        # ── [8] Timestamps are rendered in IST ───────────────────────────
        print("\n[8] Timestamps")
        probe = await conn.fetchrow(f"""
            SELECT id, created_at AS stored,
                   created_at AT TIME ZONE 'Asia/Kolkata' AS ist
              FROM new_stock_entries WHERE {COUNTED}
             ORDER BY created_at DESC LIMIT 1""")
        hit = next((r for r in rows if r["id"] == probe["id"]), None)
        check("the newest row is in the export", hit is not None)
        if hit:
            check("its timestamp is the IST instant", hit["created_at"] == probe["ist"],
                  "got %s want %s" % (hit["created_at"], probe["ist"]))
            check("and not the raw stored UTC value",
                  hit["created_at"] != probe["stored"])
            print("        stored(UTC) %s -> sheet %r"
                  % (probe["stored"], export_xlsx._dt(hit["created_at"])))
    finally:
        await conn.close()

    print("\n=== %d passed, %d failed ===\n" % (_passed, _failed))
    raise SystemExit(1 if _failed else 0)


if __name__ == "__main__":
    asyncio.run(main())
