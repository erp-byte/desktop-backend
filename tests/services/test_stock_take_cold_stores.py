"""The cold stores on Stock Take: Savla D-39 / D-514, Rishi, Eskimo, Supreme.

Their counts were loaded on 16 Sep 2026 with the warehouse spelled 'D-39' and
'D-514', while grants, the ledger and every filter value say 'D39'. An exact
match therefore dropped 387,773 kg of Savla stock from "all warehouses" and from
the D39 filter without an error. These pin the fix: one warehouse code whatever
the spelling, on the column side as well as the value side, plus the current
stock download that keeps every warehouse and floor apart.

The figures themselves are proven against the live table by
test_stock_take_latest_stock_live_sql.py and test_stock_take_floor_stock_live_sql.py.
"""
from __future__ import annotations

import asyncio
from datetime import date
from io import BytesIO

import openpyxl
import pytest
from fastapi import HTTPException

from app.modules.stock_take import floors as F
from app.modules.stock_take import router as R
from app.modules.stock_take.services import export_xlsx
from app.modules.stock_take.services import latest_stock_service as svc

HYPHEN_BLIND = "REPLACE(UPPER(BTRIM(warehouse)), '-', '')"


class FakeUser:
    def __init__(self, warehouses, floors):
        self.allowed_warehouses, self.allowed_floors, self.is_admin = warehouses, floors, True


# ── One code per warehouse ─────────────────────────────────────────────────

@pytest.mark.parametrize("given", ["D-39", "d39", " D39 ", "D39"])
def test_every_spelling_of_a_warehouse_filters_on_one_code(given):
    conds, params, applied = svc._build_filters(warehouse=[given])
    assert f"{HYPHEN_BLIND} = ANY($1::text[])" in conds
    assert params == [["D39"]]
    assert applied["warehouse"] == [given.strip()], "the echo keeps the spelling asked for"


def test_the_ledger_filter_is_hyphen_blind_too():
    conds, params = svc._build_txn_filters(start_index=3, warehouse=["D-514"], floor_name=["Savla Bond"])
    assert conds[0] == f"{HYPHEN_BLIND} = ANY($4::text[])" and params[0] == ["D514"]
    assert conds[1] == "UPPER(BTRIM(location)) = ANY($5::text[])" and params[1] == ["SAVLA BOND"]


def test_counts_and_postings_meet_at_the_same_place_key():
    """A 'D-39' count and a 'D39' posting must be one place, or the posting is
    netted against no count at all."""
    sql = " ".join(svc._stock_sql(as_of=None, adjusted_only=False, by_place=False,
                                  filters={})["ctes"].split())
    assert f"COALESCE({HYPHEN_BLIND}, '') AS k_wh" in sql
    assert f"b.k_wh = COALESCE({HYPHEN_BLIND}, '')" in sql


def test_the_page_sums_over_places_and_the_download_keeps_them_apart():
    page = " ".join(svc._stock_sql(as_of=None, adjusted_only=False, by_place=False,
                                   filters={})["ctes"].split())
    dl = " ".join(svc._stock_sql(as_of=None, adjusted_only=False, by_place=True,
                                 filters={})["ctes"].split())
    assert "AND c.k_wh = t.k_wh AND c.k_fl = t.k_fl" not in page
    assert "AND c.k_wh = t.k_wh AND c.k_fl = t.k_fl" in dl
    assert page.count("GROUP BY 1, 2, 3, 4 ") == 0
    assert dl.count("GROUP BY 1, 2, 3, 4 ") == 2, "counted and txn both keep the place"


def test_an_unusable_as_of_is_refused_not_ignored():
    with pytest.raises(ValueError):
        svc._stock_sql(as_of="13-09-2026", adjusted_only=False, by_place=True, filters={})


# ── What the dropdowns offer ───────────────────────────────────────────────

class _OptionsConn:
    """Answers fetch_places / fetch_filter_options; records the SQL."""

    def __init__(self):
        self.sql: list[str] = []

    async def fetch(self, sql, *args):
        self.sql.append(" ".join(sql.split()))
        return []


def test_places_and_options_are_keyed_by_the_code_grants_use():
    conn = _OptionsConn()
    asyncio.run(svc.fetch_places(conn))
    asyncio.run(svc.fetch_filter_options(conn))
    assert all(f"SELECT DISTINCT {HYPHEN_BLIND} AS" in s for s in conn.sql), conn.sql


# Shaped like fetch_places() now returns them on live data.
PLACES = {"A185": ["A185 Cold"], "D39": ["Savla"], "D514": ["Savla", "Savla Bond"],
          "ESKIMO": ["Eskimo"], "RISHI": ["Rishi"], "W202": ["First Floor"]}
ADMIN = FakeUser(["W202", "A185", "A68", "A101", "F53", "D-39", "D-514", "Rishi", "Supreme", "Eskimo"], None)


def test_a_hyphenated_grant_reaches_the_cold_store_floors():
    whs, by_wh = R._read_scope(ADMIN, PLACES)
    assert {"D39", "D514", "ESKIMO", "RISHI", "SUPREME"} <= set(whs)
    assert by_wh["D39"] == ["Savla"]
    assert by_wh["D514"] == ["Savla", "Savla Bond"]


def test_all_warehouses_includes_the_cold_stores():
    wh, fl, scope = R._clamp_to_read_scope(ADMIN, PLACES, None, None)
    assert {"D39", "D514", "ESKIMO", "RISHI"} <= set(wh)
    assert fl is None and scope is None, "no floor grants means every floor"


# The local (Supabase) admin on 2026-09-18, with the cold stores added to its
# warehouses: factory floor grants that name no cold-store floor.
FACTORY_FLOORS = ["Lower Basement", "Second Floor", "Second Floor Mezz", "Upper Basement",
                  "First Floor", "Terrace", "First Floor Mezz", "Sorting Area", "Printing Area",
                  "Cheese Floor", "FG store", "Dmart Production Area", "FFS Packing Area",
                  "Dmart Packing Area", "Mezzanine", "Roasting Area"]
LOCAL_ADMIN = FakeUser(["W202", "A185", "A68", "D-39", "D-514", "Rishi", "Supreme", "Eskimo"],
                       FACTORY_FLOORS)


def test_factory_floor_grants_do_not_hide_the_cold_stores():
    wh, fl, scope = R._clamp_to_read_scope(LOCAL_ADMIN, PLACES, None, None)
    assert fl is None, "the floor filter stays the caller's own choice"
    assert {"D39", "D514", "ESKIMO", "RISHI", "SUPREME", "A68"} <= set(scope["whole"])
    assert "W202" not in scope["whole"] and "A185" not in scope["whole"]
    assert "W202|FIRST FLOOR" in scope["pairs"]
    assert "A185|A185 COLD" not in scope["pairs"], "A185 Cold was not granted"
    _, by_wh = R._read_scope(LOCAL_ADMIN, PLACES)
    assert by_wh["D514"] == ["Savla", "Savla Bond"]


def test_a_floor_granted_in_one_warehouse_opens_nothing_in_another():
    """The old single floor list let an A185 grant for 'Mezzanine' read W202's
    rows on a floor of the same name; pairs keep each grant in its building."""
    places = {"A185": ["Mezzanine"], "W202": ["Mezzanine", "Terrace"]}
    _, _, scope = R._clamp_to_read_scope(FakeUser(["A185", "W202"], ["Mezzanine"]), places, None, None)
    assert scope == {"whole": [], "pairs": ["A185|MEZZANINE", "W202|MEZZANINE"]}
    _, _, scope = R._clamp_to_read_scope(FakeUser(["A185", "W202"], ["Terrace"]), places, None, None)
    assert scope["pairs"] == ["W202|TERRACE"]


def test_a_cold_store_floor_can_be_picked_by_a_floor_restricted_caller():
    wh, fl, scope = R._clamp_to_read_scope(LOCAL_ADMIN, PLACES, ["D39"], ["Savla Bond"])
    assert wh == ["D39"] and fl == ["Savla Bond"] and "D514" in scope["whole"]
    with pytest.raises(HTTPException) as exc:
        R._clamp_to_read_scope(LOCAL_ADMIN, PLACES, None, ["A185 Cold"])
    assert exc.value.detail["error"] == "floor_not_allowed"


def test_the_place_scope_is_one_predicate_on_both_tables_with_the_right_numbers():
    scope = {"whole": ["D39"], "pairs": ["W202|TERRACE"]}
    conds, params, applied = svc._build_filters(warehouse=["D39", "W202"], place_scope=scope)
    want = (f"({HYPHEN_BLIND} = ANY($2::text[])"
            f" OR {HYPHEN_BLIND} || '|' || UPPER(BTRIM(floor_name)) = ANY($3::text[]))")
    assert want in conds and params[1:3] == [["D39"], ["W202|TERRACE"]]
    assert "place_scope" not in applied and "placeScope" not in applied
    tconds, tparams = svc._build_txn_filters(start_index=5, warehouse=["D39"], place_scope=scope)
    assert (f"({HYPHEN_BLIND} = ANY($7::text[])"
            f" OR {HYPHEN_BLIND} || '|' || UPPER(BTRIM(location)) = ANY($8::text[]))") in tconds
    assert tparams[1:3] == [["D39"], ["W202|TERRACE"]]


def test_the_download_refuses_a_warehouse_outside_the_profile():
    with pytest.raises(HTTPException) as exc:
        R._clamp_to_read_scope(FakeUser(["W202"], None), PLACES, ["D-39"], None)
    assert exc.value.status_code == 403
    assert exc.value.detail["error"] == "warehouse_not_allowed"


@pytest.mark.parametrize("code, label", [
    ("D39", "Savla D-39"), ("D-514", "Savla D-514"), ("ESKIMO", "Eskimo"),
    ("Rishi", "Rishi"), ("W202", "W202"), ("", ""),
])
def test_warehouse_names(code, label):
    assert F.warehouse_label(code) == label


def test_filter_options_name_and_group_the_cold_stores(monkeypatch):
    async def places(conn):
        return PLACES

    async def options(conn):
        return {"warehouses": sorted(PLACES), "floors": [], "item_types": [], "stock_types": []}

    class _Pool:
        def acquire(self):
            class _Ctx:
                async def __aenter__(self_inner):
                    return object()

                async def __aexit__(self_inner, *a):
                    return False
            return _Ctx()

    class _Req:
        class app:
            class state:
                db_pool = _Pool()

    monkeypatch.setattr(svc, "fetch_places", places)
    monkeypatch.setattr(svc, "fetch_filter_options", options)
    out = asyncio.run(R.filter_options(_Req(), user=ADMIN))
    assert out["warehouse_labels"]["D39"] == "Savla D-39"
    assert out["warehouse_labels"]["W202"] == "W202"
    assert set(out["cold_warehouses"]) == {"D39", "D514", "ESKIMO", "RISHI", "SUPREME"}
    assert "W202" not in out["cold_warehouses"]


# ── The download workbook ──────────────────────────────────────────────────

def _row(wh, floor_key, floor, item, kg, *, counted=0.0, stock="Fresh Stock", day=None):
    return {"warehouse": wh, "floor_key": floor_key, "floor": floor, "item_name": item,
            "item_type": "RM", "item_category": "Seeds", "item_subcategory": "",
            "stock_type": stock, "total_quantity": 1.0, "counted_weight": counted,
            "net_adjustment_kg": kg - counted, "total_weight": kg, "entry_count": 1,
            "transaction_count": 0, "last_counted_date": day, "days_since_count": None}


ROWS = [
    _row("W202", "STORE", "STORE", "Almonds", 10.0),
    _row("W202", "STORE", "Store", "Cashew", 5.0),           # same floor, other spelling
    _row("W202", "FIRST FLOOR", "First Floor", "Pista", 2.5),
    _row("D39", "SAVLA", "Savla", "Sunflower Seeds", 52550.0, counted=52550.0, day=date(2026, 9, 13)),
    _row("D39", "SAVLA", "Savla", "Cashew", 100.0, counted=100.0, stock="Off Grade/Rejection",
         day=date(2026, 9, 13)),
    _row("ESKIMO", "ESKIMO", "Eskimo", "Dates", 243800.0, counted=243800.0, day=date(2026, 9, 13)),
]


def _summary(rows, filters=None):
    wb = openpyxl.load_workbook(BytesIO(export_xlsx.build_stock_workbook(rows, filters or {}, "Tester").read()))
    ws = wb["Summary"]
    lines = [r for r in ws.iter_rows(min_row=7, values_only=True) if r[0]]
    return wb, lines


def test_summary_has_every_floor_a_total_per_warehouse_and_the_split():
    wb, lines = _summary(ROWS)
    firsts = [(r[0], r[1]) for r in lines]
    assert firsts == [
        ("W202", "First Floor"), ("W202", "Store"), ("W202 total", None),
        ("Savla D-39", "Savla"), ("Savla D-39 total", None),
        ("Eskimo", "Eskimo"), ("Eskimo total", None),
        ("Factories and godowns", None), ("Cold storage", None), ("GRAND TOTAL", None),
    ], "factories first in walking order, then the cold stores, STORE/Store as one floor"
    by = {r[0] if r[1] is None else (r[0], r[1]): r for r in lines}
    assert by[("W202", "Store")][2] == 2 and by[("W202", "Store")][5] == 15.0
    assert by["Savla D-39 total"][5] == 52650.0 and by["Savla D-39 total"][6] == 100.0
    assert by["Cold storage"][5] == 52650.0 + 243800.0
    assert by["GRAND TOTAL"][5] == sum(r["total_weight"] for r in ROWS)
    assert by["GRAND TOTAL"][2] == len(ROWS)
    detail = list(wb["Stock by floor"].iter_rows(min_row=2, values_only=True))
    assert len(detail) == len(ROWS)
    assert {d[1] for d in detail if d[0] == "W202"} == {"Store", "First Floor"}
    assert {d[0] for d in detail} == {"W202", "Savla D-39", "Eskimo"}


def test_one_kind_only_has_no_split_lines():
    _, lines = _summary([r for r in ROWS if r["warehouse"] == "W202"])
    assert [r[0] for r in lines][-1] == "GRAND TOTAL"
    assert "Cold storage" not in [r[0] for r in lines]


def test_an_empty_download_says_so():
    wb = openpyxl.load_workbook(BytesIO(export_xlsx.build_stock_workbook([], {}, "T").read()))
    values = [c for row in wb["Summary"].iter_rows(values_only=True) for c in row if c]
    assert "No stock matched these filters." in values
    assert wb["Stock by floor"].max_row == 1


def test_the_header_states_the_filters_by_name():
    wb, _ = _summary(ROWS, {"warehouse": ["D39"], "stockType": ["Fresh Stock"]})
    assert wb["Summary"]["A3"].value == "Warehouse: Savla D-39; Stock type: Fresh Stock"
    wb, _ = _summary(ROWS, {})
    assert wb["Summary"]["A3"].value == "Warehouse: all"


@pytest.mark.parametrize("term", ["D-39", "D39", "d-39"])
def test_searching_a_warehouse_is_hyphen_blind_on_both_sides(term):
    """Savla's counts say 'D-39' and its postings 'D39'; a search for either
    spelling must find both, or the figure nets a count against nothing."""
    conds, params, _ = svc._build_filters(search=term)
    tconds, tparams = svc._build_txn_filters(start_index=0, search=term)
    want = "REPLACE(UPPER(COALESCE(warehouse, '')), '-', '') LIKE REPLACE($1, '-', '')"
    assert any(want in c for c in conds) and any(want in c for c in tconds)
    assert params[-1] == tparams[-1] == "%" + term.upper() + "%"


def test_warehouse_totals_count_an_item_once_however_many_floors_hold_it():
    """The screen's Items figure counts item + stock type once; so must the totals."""
    rows = [
        {**_row("W202", "STORE", "Store", "Almonds", 10.0), "item_key": "ALMONDS"},
        {**_row("W202", "FIRST FLOOR", "First Floor", "Almonds ", 4.0), "item_key": "ALMONDS"},
        {**_row("W202", "FIRST FLOOR", "First Floor", "Almonds", 1.0, stock="Off Grade/Rejection"),
         "item_key": "ALMONDS"},
    ]
    _, lines = _summary(rows)
    by = {r[0] if r[1] is None else (r[0], r[1]): r for r in lines}
    assert by[("W202", "First Floor")][2] == 2
    assert by["W202 total"][2] == 2, "Almonds fresh + Almonds off grade"
    assert by["GRAND TOTAL"][2] == 2 and by["GRAND TOTAL"][5] == 15.0


# ── Narrow or whole is decided by the grants, never by today's rows ────────

COLD_PLACES = {"D39": ["Savla"], "D514": ["Savla", "Savla Bond"], "W202": ["Store"], "A68": ["Store", "Bay 1"]}


def test_a_cold_floor_grant_set_in_the_table_narrows_every_cold_store():
    """[D-39, D-514] + [Savla Bond] means bonded stock only, not all of D-39."""
    _, by_wh, whole = R._read_scope_detail(FakeUser(["D-39", "D-514"], ["Savla Bond"]), COLD_PLACES)
    assert whole == [] and by_wh["D39"] == [] and by_wh["D514"] == ["Savla Bond"]


def test_a_misspelt_cold_floor_grant_shows_nothing_rather_than_everything():
    _, by_wh, whole = R._read_scope_detail(FakeUser(["D-514"], ["Savla-Bond"]), COLD_PLACES)
    assert whole == [] and by_wh["D514"] == []


def test_deleting_counts_cannot_widen_a_grant():
    """If D-514's 'Savla' rows go (a re-sync, a relabel), [Savla] still narrows it."""
    places = {"D39": ["Savla"], "D514": ["Savla Bond"]}
    _, by_wh, whole = R._read_scope_detail(FakeUser(["D-39", "D-514"], ["Savla"]), places)
    assert whole == [] and by_wh["D514"] == []


def test_a_godown_sharing_a_floor_name_with_w202_stays_whole():
    """'Store' picked for W202 on the admin screen says nothing about A68."""
    _, by_wh, whole = R._read_scope_detail(FakeUser(["W202", "A68"], ["Store"]), COLD_PLACES)
    assert whole == ["A68"] and by_wh["A68"] == ["Store", "Bay 1"] and by_wh["W202"] == ["Store"]
